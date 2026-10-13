"""Document parsing tests: TXT/PDF/DOCX extraction and failure modes."""

from app.ingestion.docparse import parse_upload

TXT = """Sam Lee
sam@example.com 555-0199

EXPERIENCE
Backend Engineer, Z 2020-2025
- Built Python APIs

EDUCATION
BS Computer Science, 2016-2020

SKILLS
Python, Docker
"""


def _pdf_bytes(lines: list[str]) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in lines:
        pdf.cell(0, 6, line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def _docx_bytes(lines: list[str]) -> bytes:
    import io

    from docx import Document

    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_txt_sections_and_contact():
    out = parse_upload("r.txt", "text/plain", TXT.encode())
    assert out["parse_status"] == "ok"
    assert "experience" in out["sections"] and "education" in out["sections"]
    assert "sam@example.com" in out["contact"]["emails"]
    assert "2020" in out["contact"]["years"]


def test_pdf_text_extraction():
    out = parse_upload("r.pdf", "application/pdf",
                       _pdf_bytes(["Sam Lee", "EXPERIENCE", "Built Python APIs"]))
    assert out["parse_status"] == "ok", out["warnings"]
    assert "Python" in out["text"]
    assert "experience" in out["sections"]


def test_docx_text_extraction():
    out = parse_upload("r.docx",
                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                       _docx_bytes(["Sam Lee", "SKILLS", "Python, Docker"]))
    assert out["parse_status"] == "ok", out["warnings"]
    assert "Docker" in out["text"]


def test_empty_document():
    out = parse_upload("empty.txt", "text/plain", b"   \n  ")
    assert out["parse_status"] == "empty"
    assert out["text"] == ""


def test_blank_pdf_reports_no_text():
    out = parse_upload("blank.pdf", "application/pdf", _pdf_bytes([]))
    assert out["parse_status"] in ("empty", "needs_ocr")
    assert out["warnings"]


def test_garbage_pdf_fails_cleanly():
    out = parse_upload("bad.pdf", "application/pdf", b"not a pdf at all %%%")
    assert out["parse_status"] == "failed"
    assert out["text"] == ""


def test_unsupported_extension():
    out = parse_upload("x.exe", "application/octet-stream", b"MZ")
    assert out["parse_status"] == "failed"


def test_upload_stores_parsed_doc(client, auth_headers):
    r = client.post("/api/resumes",
                    files={"file": ("r.pdf", _pdf_bytes(["Sam Lee", "EXPERIENCE", "Built Python APIs"]),
                                    "application/pdf")},
                    headers=auth_headers)
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["parse_status"] == "parsed"
    got = client.get(f"/api/resumes/{body['id']}", headers=auth_headers).json()
    assert "Python" in got["raw_text"]
